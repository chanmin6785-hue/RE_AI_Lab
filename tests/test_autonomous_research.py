import json
import copy
import re
from pathlib import Path
import shutil
import unittest
from datetime import datetime, timedelta, timezone
from uuid import uuid4
from zoneinfo import ZoneInfo

from pptx import Presentation
from docx import Document

from app import Lab, next_credit_check
from baze_models import BazeError, Reply
from research_selection import assign_roles, select_topic
from research_workflow import simulation_protocol


class ResearchClient:
    calls = []
    balance = 8000
    monthly_balance = None
    selected_domain = "AI"

    def __init__(self, key, timeout):
        self.api_key = key

    def list_models(self):
        return ["gpt-test", "claude-test", "gemini-test", "grok-test"]

    def credits(self):
        amount = type(self).balance
        monthly = type(self).monthly_balance
        return {"monthly_allocated": {"remaining": amount if monthly is None else monthly},
                "total": {"remaining": amount}}

    def ask(self, provider, prompt, model, max_tokens):
        type(self).calls.append((provider, prompt, model))
        phase = prompt.split("단계: ", 1)[1].split(";", 1)[0] if "단계: " in prompt else ""
        text = {
            "topic_proposal": lambda: "CANDIDATES_JSON: " + json.dumps([
                {"title": f"{provider}의 {re.search(r'첫 후보는 (AI|물리|수학|정보) 분야', prompt).group(1)} 분야 연구 후보",
                 "question": "독립 변수 변화가 측정 가능한 결과에 어떤 영향을 주는가?",
                 "domain": re.search(r'첫 후보는 (AI|물리|수학|정보) 분야', prompt).group(1),
                 "method": "공개 자료와 반복 가능한 수식 또는 모델 호출로 비교한다"},
                {"title": f"{provider}의 정보 융합 연구 후보",
                 "question": "모형과 기준선 사이의 수치 차이는 재현되는가?",
                 "domain": "정보" if re.search(r'첫 후보는 (AI|물리|수학|정보) 분야', prompt).group(1) != "정보" else "AI",
                 "method": "동일한 입력 격자에서 두 모형 결과를 비교한다"}], ensure_ascii=False),
            "topic_critique": lambda: "두 후보의 자료 접근성과 대조군을 검토했다.\nCANDIDATES_JSON: " + json.dumps([
                {"title": f"{provider}의 개선된 수학 연구 후보",
                 "question": "수식의 매개변수 변화가 예측 오차에 영향을 주는가?",
                 "domain": "수학", "method": "고정된 입력 격자와 기준선 수식을 반복 비교한다"}], ensure_ascii=False),
            "topic_ballot": lambda: "동일한 기준으로 후보를 비교했다.\nRANKING: " + ",".join(
                ([next((cid for cid, domain in re.findall(r"(C\d{2}) \[(AI|물리|수학|정보);", prompt)
                        if domain == type(self).selected_domain), "C01")] +
                 [cid for cid in re.findall(r"(C\d{2}) \[", prompt)
                  if cid != next((cid for cid, domain in re.findall(r"(C\d{2}) \[(AI|물리|수학|정보);", prompt)
                                  if domain == type(self).selected_domain), "C01")])[:3]),
            "role_trial": lambda: "이론: 가설을 수식으로 검증한다. 실험: 통제군을 둔다. 구현: 재현 코드를 설계한다. 비판: 반례를 검토한다.",
            "role_review": lambda: "실제 작업 표본을 평가했다.\nSCORES_JSON: " + json.dumps({
                name: {competency: 5 if index == position else 2
                       for position, competency in enumerate(("theory", "experiment", "implementation", "critique"))}
                for index, name in enumerate(("gpt", "claude", "gemini", "grok"))}),
            "simulation_design": lambda: '가정에 따른 수식 계산이다.\nSIMULATION_JSON: {"expression":"x*x","baseline":"x","x_start":0,"x_stop":10,"points":30,"parameters":{},"hypothesis":"차이가 증가한다","limitations":"실측이 아니다"}',
            "role_selection": "MY_ROLE: 실험 설계\n" + "\n".join(
                f"ROLE_{name.upper()}: 검증과 분석" for name in ("gpt", "claude", "gemini", "grok")),
            "topic_selection": "RESEARCH_TOPIC: 네 모델 답변 비교\nRESEARCH_QUESTION: 동일 질문에 대한 응답은 어떻게 다른가?",
            "plan_discussion": "대조군과 재현 방법을 검토한다.\nNEXT: DRAFT",
            "plan_report": "# 연구계획보고서\n\nAI 생성 사례를 쓰고 실제 모델 응답을 관찰한다.",
            "plan_slides": "SLIDE 1: 연구 질문\n- 동일 질문의 응답 차이\nSLIDE 2: 실험 계획\n- 네 모델 독립 호출",
            "plan_script": "# 계획발표 대본\n첫 장은 연구 질문입니다.",
            "experiment_discussion": "대조군을 유지한다.\nNEXT: CASES" if "실험 주기: 1" in prompt and "실험 분석 자료: 아직" in prompt else (
                "결과를 중간 발표에서 검토한다.\nNEXT: MIDDLE" if "실험 주기: 1" in prompt else (
                    "새 사례로 재검증한다.\nNEXT: CASES" if "실험 분석 자료: 아직" in prompt else "최종 검토로 간다.\nNEXT: FINAL")),
            "experiment_cases": 'CASES_JSON: [{"prompt":"2+2는?", "expected":"4", "source_note":"AI 생성"}]',
            "experiment_analysis": "# 실제 응답 해석\n문자 수는 품질이 아니다.",
            "middle_slides": "SLIDE 1: 중간 결과\n- 실제 모델 응답을 수집했다.",
            "middle_script": "# 중간발표 대본\n관찰값을 설명합니다.",
            "final_discussion": "증거와 한계를 확인했다.\nNEXT: FINAL",
            "final_report": "# 최종 연구보고서\n실제 모델 응답과 AI 생성 사례를 구분한다.",
            "final_slides": "SLIDE 1: 최종 결과\n- 관찰값과 한계",
            "final_script": "# 최종발표 대본\n결과와 한계를 발표합니다.",
        }.get(phase, "4")
        if callable(text):
            text = text()
        return Reply(provider, model, text, {"input_tokens": 10}, {"id": f"fake-{len(type(self).calls)}"})


class AutonomousResearchTest(unittest.TestCase):
    def setUp(self):
        self.temp = Path(__file__).resolve().parent / ("auto_test_" + uuid4().hex)
        self.temp.mkdir()
        ResearchClient.calls = []
        ResearchClient.balance = 8000
        ResearchClient.monthly_balance = None
        ResearchClient.selected_domain = "AI"
        self.lab = Lab(ResearchClient, self.temp)
        for actor in ("user", "friend", "member3", "member4"):
            self.lab.connect("KEY_" + actor, actor=actor)
        self.models = {name: name + "-test" for name in ("gpt", "claude", "gemini", "grok")}
        self.assignments = dict(zip(self.models, ("user", "friend", "member3", "member4")))

    def tearDown(self):
        if self.lab.worker:
            self.lab.worker.join(5)
        shutil.rmtree(self.temp)

    def advance_until(self, status, limit=40):
        for _ in range(limit):
            if self.lab.worker:
                self.lab.worker.join(5)
            if self.lab.status == status:
                return
            self.lab.auto["last_wake_at"] = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()
            self.lab.auto_wake()
        self.fail(f"did not reach {status}: {self.lab.status} {self.lab.auto.get('stage')} {self.lab.error}")

    def test_plan_middle_final_revisions_real_evidence_and_pptx(self):
        self.lab.start(None, self.models, self.assignments)
        self.advance_until("waiting_review")
        self.assertEqual(self.lab.auto["stage"], "plan_review")
        self.assertEqual(self.lab.auto["domain"], "AI")
        self.assertEqual({item["provider"] for item in self.lab.auto["candidates"]}, set(self.models))
        self.assertEqual({item["domain"] for item in self.lab.auto["candidates"]}, {"AI", "물리", "수학", "정보"})
        self.assertEqual(len(self.lab.auto["ballots"]), 4)
        self.assertEqual(len(set(self.lab.roles.values())), 4)
        self.assertIn("작업 표본", (self.lab.project / "research/role_selection.md").read_text(encoding="utf-8"))
        self.assertIn("후보 비교", (self.lab.project / "research/topic_selection.md").read_text(encoding="utf-8"))
        self.assertEqual(len(self.lab.deliverables()), 8)
        report = self.lab.project / "presentations" / "plan" / "report.docx"
        self.assertIn("연구계획보고서", " ".join(p.text for p in Document(str(report)).paragraphs))
        plan = self.lab.project / "presentations" / "plan" / "slides.pptx"
        self.assertGreaterEqual(len(Presentation(str(plan)).slides), 3)
        self.assertEqual(len(self.lab.auto["observations"]), 0)
        self.lab.review("plan_review", False, "대조 조건을 더 명확히", "friend")
        self.advance_until("waiting_review")
        self.assertEqual(self.lab.auto["stage"], "plan_review")
        self.assertEqual(self.lab.auto["revision"]["plan"], 2)
        self.lab.review("plan_review", True, "계획 확인", "user")
        self.advance_until("waiting_review")
        self.assertEqual(self.lab.auto["stage"], "middle_review")
        self.assertGreater(len(self.lab.auto["observations"]), 0)
        self.assertTrue((self.lab.project / "data/processed/observations.jsonl").exists())
        cases = (self.lab.project / "datasets/cycle_01_ai_generated_cases.jsonl").read_text(encoding="utf-8")
        self.assertIn("AI_GENERATED", cases)
        analysis = json.loads((self.lab.project / "results/cycle_01_analysis.json").read_text(encoding="utf-8"))
        self.assertEqual(analysis["evidence"], "OBSERVED_BAZE")
        self.lab.review("middle_review", True, "중간 검토 완료", "member3")
        self.advance_until("waiting_review")
        self.assertEqual(self.lab.auto["stage"], "final_review")
        self.assertTrue((self.lab.project / "presentations/final/report.md").exists())
        self.lab.review("final_review", True, "최종 검토 완료", "member4")
        self.assertEqual(self.lab.status, "completed")
        self.assertEqual([row["stage"] for row in self.lab.auto["reviews"]],
                         ["plan_review", "plan_review", "middle_review", "final_review"])

    def test_physics_topic_uses_simulation_not_model_comparison(self):
        ResearchClient.selected_domain = "물리"
        self.lab.start(None, self.models, self.assignments)
        self.advance_until("waiting_review")
        self.assertEqual(self.lab.auto["domain"], "물리")
        self.assertIn("물리", self.lab.topic)
        self.lab.review("plan_review", True, "수식 모델 검토", "user")
        self.advance_until("waiting_review")
        self.assertEqual(self.lab.auto["stage"], "middle_review")
        self.assertEqual(self.lab.auto["observations"], [])
        self.assertFalse((self.lab.project / "data/processed/observations.jsonl").exists())
        analysis = json.loads((self.lab.project / "results/cycle_01_analysis.json").read_text(encoding="utf-8"))
        self.assertEqual(analysis["evidence"], "SIMULATED_MODEL")
        self.assertEqual(analysis["n"], 30)
        dataset = (self.lab.project / "datasets/cycle_01_simulated_data.csv").read_text(encoding="utf-8-sig")
        self.assertIn("SIMULATED_MODEL", dataset)

    def test_restart_preserves_prior_project_and_starts_new_selection(self):
        self.lab.start(None, self.models, self.assignments)
        self.advance_until("waiting_review")
        prior = self.lab.project
        prior_turns = self.lab.turn_no
        self.lab.restart()
        self.lab.worker.join(5)
        self.assertNotEqual(self.lab.project, prior)
        self.assertEqual(self.lab.auto["selection_version"], 2)
        archived = json.loads((prior / "research/archived_state_before_restart.json").read_text(encoding="utf-8"))
        self.assertEqual(archived["turn_no"], prior_turns)
        self.assertEqual(archived["status"], "waiting_review")
        self.assertTrue((prior / "presentations/plan/report.md").exists())

    def test_selection_counts_all_four_ballots_and_excludes_self_role_scores(self):
        candidates = [{"id": "C01", "title": "첫 후보"}, {"id": "C02", "title": "둘째 후보"},
                      {"id": "C03", "title": "셋째 후보"}]
        ballots = {"gpt": ["C01", "C02", "C03"], "claude": ["C02", "C03", "C01"],
                   "gemini": ["C02", "C03", "C01"], "grok": ["C03", "C02", "C01"]}
        winner, comparison = select_topic(candidates, ballots)
        self.assertEqual(winner["id"], "C02")
        self.assertEqual(len(comparison), 3)
        competences = ("theory", "experiment", "implementation", "critique")
        providers = ("gpt", "claude", "gemini", "grok")
        reviews = {judge: {provider: {task: (5 if task == competences[index] else 2)
                                      for task in competences}
                           for index, provider in enumerate(providers)} for judge in providers}
        reviews["gpt"]["gpt"] = dict.fromkeys(competences, 1)
        roles, scores = assign_roles(providers, reviews)
        self.assertEqual(roles["gpt"], "이론·수학적 검증")
        self.assertEqual(scores["gpt"]["theory"], 5)

    def test_simulation_protocol_rejects_model_authored_code(self):
        malicious = {"expression": "__import__('os').system('dir')", "baseline": "x",
                     "x_start": 0, "x_stop": 1, "points": 20}
        self.assertIsNone(simulation_protocol("SIMULATION_JSON: " + json.dumps(malicious)))

    def test_low_credit_waits_without_paid_call(self):
        ResearchClient.balance = 0
        self.lab.start(None, self.models, self.assignments)
        self.lab.worker.join(5)
        self.assertEqual(self.lab.status, "waiting_budget")
        self.assertEqual(ResearchClient.calls, [])
        self.assertTrue(self.lab.auto["next_due_at"])
        self.lab.auto_wake()
        self.assertEqual(ResearchClient.calls, [])
        ResearchClient.balance = 8000
        self.lab.auto["next_due_at"] = (datetime.now(timezone.utc) - timedelta(minutes=1)).isoformat()
        self.lab.auto["last_wake_at"] = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()
        self.lab.auto_wake()
        self.lab.worker.join(5)
        self.assertGreater(len(ResearchClient.calls), 0)

    def test_next_credit_check_uses_upcoming_midnight_window(self):
        kst = ZoneInfo("Asia/Seoul")
        day = datetime(2026, 10, 14, 0, 5, tzinfo=kst)
        self.assertEqual(next_credit_check(day), datetime(2026, 10, 14, 0, 10, tzinfo=kst))
        self.assertEqual(next_credit_check(day + timedelta(minutes=10)),
                         datetime(2026, 10, 15, 0, 10, tzinfo=kst))

    def test_positive_total_credit_allows_calls_even_when_monthly_is_zero(self):
        ResearchClient.balance = 1
        ResearchClient.monthly_balance = 0
        self.lab.start(None, self.models, self.assignments)
        self.lab.worker.join(5)
        self.assertGreater(len(ResearchClient.calls), 0)
        self.assertNotEqual(self.lab.status, "waiting_budget")

    def test_gateway_credit_rejection_waits_then_retries_with_new_audit_number(self):
        ResearchClient.balance = 1
        for client in self.lab.clients.values():
            client.ask = lambda *args, **kwargs: (_ for _ in ()).throw(BazeError("HTTP 402: insufficient credits"))
        self.lab.start(None, self.models, self.assignments)
        self.lab.worker.join(5)
        self.assertEqual(self.lab.status, "waiting_budget")
        self.assertEqual(self.lab.activities[-1]["status"], "failed")
        failed_turn = self.lab.activities[-1]["turn"]
        for client in self.lab.clients.values():
            del client.ask
        self.lab.auto["next_due_at"] = (datetime.now(timezone.utc) - timedelta(minutes=1)).isoformat()
        self.lab.auto["last_wake_at"] = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()
        self.lab.auto_wake()
        self.lab.worker.join(5)
        self.assertGreater(len(ResearchClient.calls), 0)
        self.assertGreater(self.lab.activities[-1]["turn"], failed_turn)

    def test_gateway_timeout_waits_and_retries_same_speaker_with_new_audit_number(self):
        for client in self.lab.clients.values():
            client.ask = lambda *args, **kwargs: (_ for _ in ()).throw(
                BazeError("연결 오류: The read operation timed out"))
        self.lab.start(None, self.models, self.assignments)
        self.lab.worker.join(5)
        self.assertEqual(self.lab.status, "waiting_retry")
        failed = self.lab.activities[-1]
        self.assertEqual(failed["status"], "failed")
        self.assertIn("중복 크레딧", self.lab.auto["reason"])
        self.lab.auto_wake()
        self.assertEqual(len(self.lab.activities), 1)
        for client in self.lab.clients.values():
            del client.ask
        self.lab.auto["next_due_at"] = (datetime.now(timezone.utc) - timedelta(minutes=1)).isoformat()
        self.lab.auto["last_wake_at"] = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()
        self.lab.auto_wake()
        self.lab.worker.join(5)
        completed = next(item for item in self.lab.activities if item["status"] == "completed")
        self.assertEqual((completed["phase"], completed["provider"]),
                         (failed["phase"], failed["provider"]))
        self.assertGreater(completed["turn"], failed["turn"])

    def test_restore_recovers_old_paused_timeout_once(self):
        self.lab.start(None, self.models, self.assignments)
        self.lab.worker.join(5)
        snapshot = copy.deepcopy(self.lab._snapshot())
        snapshot["status"] = "paused"
        snapshot["error"] = "연결 오류: The read operation timed out"
        snapshot["auto"]["stage"] = "topic_critiques"
        snapshot["activities"].append({"turn": snapshot["turn_no"] + 1,
                                       "phase": "topic_critique", "provider": "grok",
                                       "status": "failed", "error": snapshot["error"]})

        class RecoveryStore:
            def __init__(self, state):
                self.state = copy.deepcopy(state)

            def load(self, _target):
                return copy.deepcopy(self.state)

            def save(self, _project, state):
                self.state = copy.deepcopy(state)

        store = RecoveryStore(snapshot)
        restored = Lab(ResearchClient, self.temp, project_store=store)
        self.assertEqual(restored.status, "running")
        self.assertEqual(restored.error, "")
        self.assertEqual(restored.auto["timeout_recovery_version"], 1)
        restored_again = Lab(ResearchClient, self.temp, project_store=store)
        self.assertEqual(sum("시간 초과로 확인하지 못한 발언" in item["content"]
                             for item in restored_again.messages), 1)

    def test_critique_prompt_contains_candidates_without_duplicate_proposal_transcripts(self):
        self.lab.start(None, self.models, self.assignments)
        self.lab.worker.join(5)
        self.assertEqual(self.lab.auto["stage"], "topic_critiques")
        for item in self.lab.auto["candidates"]:
            item["method"] = "독립변수와 대조군을 반복 측정한다. " * 60
        self.lab.messages.append({"actor": "model", "provider": "gemini",
                                  "phase": "topic_proposal", "content": "DUPLICATE_PROPOSAL " * 300})
        prompt = self.lab._autonomous_prompt("topic_critique", "grok", 0)
        self.assertLess(len(prompt.encode("utf-8")), 20000)
        self.assertNotIn("DUPLICATE_PROPOSAL", prompt)
        for item in self.lab.auto["candidates"]:
            self.assertIn(item["id"], prompt)

    def test_overview_reads_complete_archived_dialogue(self):
        self.lab.start(None, self.models, self.assignments)
        self.advance_until("waiting_review")
        with self.lab.lock:
            for number in range(205):
                self.lab._append("system", f"기록 확인 {number}")
        overview = self.lab.research_overview()
        self.assertGreater(overview["counts"]["all"], len(self.lab.messages))
        self.assertEqual(overview["counts"]["models"], len(ResearchClient.calls))
        self.assertTrue(any(section["key"] == "roles" for section in overview["sections"]))
        self.assertEqual(len(overview["roles"]), 4)
        self.assertIn("검토", overview["current"]["task"])

    def test_restore_recreates_empty_folders_and_missing_role_file(self):
        self.lab.start(None, self.models, self.assignments)
        self.advance_until("waiting_review")
        self.assertEqual(self.lab.auto["stage"], "plan_review")
        snapshot = copy.deepcopy(self.lab._snapshot())
        snapshot["status"] = "waiting_budget"
        snapshot["error"] = "[Errno 2] No such file or directory: 'research/role_selection.md'"
        files = {path.relative_to(self.lab.project): path.read_bytes()
                 for path in self.lab.project.rglob("*") if path.is_file()
                 and path.name != "role_selection.md"}

        class MissingRoleStore:
            def load(self, target):
                project = target / snapshot["project_name"]
                for relative, content in files.items():
                    path = project / relative
                    path.parent.mkdir(parents=True, exist_ok=True)
                    path.write_bytes(content)
                return copy.deepcopy(snapshot)

            def save(self, project, state):
                files.update({path.relative_to(project): path.read_bytes()
                              for path in project.rglob("*") if path.is_file()})

        restored = Lab(ResearchClient, self.temp / "restored", project_store=MissingRoleStore())
        role_file = restored.project / "research" / "role_selection.md"
        self.assertTrue(role_file.exists())
        self.assertIn("GPT", role_file.read_text(encoding="utf-8"))
        self.assertIn(Path("research/role_selection.md"), files)
        self.assertEqual(restored.error, "")
        self.assertTrue((restored.project / "presentations" / "final").is_dir())

    def test_restore_recovers_claude_output_limit_once(self):
        self.lab.start(None, self.models, self.assignments)
        self.lab.worker.join(5)
        snapshot = copy.deepcopy(self.lab._snapshot())
        snapshot["status"] = "paused"
        snapshot["error"] = ('텍스트 답변이 없습니다. 토큰 한도·모델 지원·필터 결과를 확인하세요. '
                             '{"stop_reason": "max_tokens"}')
        snapshot["auto"]["stage"] = "plan_discussion"
        snapshot["auto"]["credit_rule_version"] = "zero_only"
        snapshot["turn_no"] += 1
        snapshot["activities"].append({"turn": snapshot["turn_no"],
                                       "phase": "plan_discussion", "provider": "claude",
                                       "status": "failed", "error": snapshot["error"]})
        snapshot["error"] = "서버가 재시작되었습니다. 기록을 복원했습니다. 재개 전에 마지막 호출을 확인하세요."

        class RecoveryStore:
            def __init__(self, state):
                self.state = copy.deepcopy(state)

            def load(self, _target):
                return copy.deepcopy(self.state)

            def save(self, _project, state):
                self.state = copy.deepcopy(state)

        store = RecoveryStore(snapshot)
        restored = Lab(ResearchClient, self.temp / "recovered", project_store=store)
        self.assertEqual(restored.status, "running")
        self.assertEqual(restored.error, "")
        self.assertEqual(restored.auto["claude_output_recovery_version"], 1)
        self.assertEqual(store.state["status"], "running")
        note = "Claude의 출력 한도 오류를 수정했습니다."
        self.assertEqual(sum(note in item["content"] for item in restored.messages), 1)
        restored_again = Lab(ResearchClient, self.temp / "recovered-again", project_store=store)
        self.assertEqual(restored_again.status, "running")
        self.assertEqual(sum(note in item["content"] for item in restored_again.messages), 1)


if __name__ == "__main__":
    unittest.main()
