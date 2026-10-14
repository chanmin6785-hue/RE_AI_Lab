"""Reproducible research artifacts and bounded empirical model experiments."""

from __future__ import annotations

import ast
import csv
import json
import math
import random
import re
import statistics
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt as DocPt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches, Pt


STAGES = ("topic_proposals", "topic_critiques", "topic_ballots", "role_trials", "role_reviews",
          "roles", "topic", "plan_discussion", "plan_report", "plan_slides",
          "plan_script", "plan_review", "experiment_discussion", "experiment_cases", "simulation_design",
          "experiment_run", "experiment_analysis", "middle_slides", "middle_script",
          "middle_review", "final_discussion", "final_report", "final_slides",
          "final_script", "final_review", "completed")
REVIEW_STAGES = {"plan_review": "연구 계획", "middle_review": "중간 발표", "final_review": "최종 결과"}


_FUNCTIONS = {"sin": math.sin, "cos": math.cos, "tan": math.tan, "sqrt": math.sqrt,
              "exp": math.exp, "log": math.log, "abs": abs}
_OPERATORS = {ast.Add: lambda a, b: a + b, ast.Sub: lambda a, b: a - b,
              ast.Mult: lambda a, b: a * b, ast.Div: lambda a, b: a / b,
              ast.Pow: lambda a, b: a ** b}


def simulation_protocol(text: str) -> dict | None:
    """Read a declarative numerical model; never execute model-authored code."""
    from research_selection import marked_json
    data = marked_json(text, "SIMULATION_JSON")
    if not isinstance(data, dict):
        return None
    try:
        expression, baseline = str(data["expression"]), str(data["baseline"])
        start, stop = float(data["x_start"]), float(data["x_stop"])
        points = int(data["points"])
        parameters = data.get("parameters", {})
        if not isinstance(parameters, dict) or len(parameters) > 8:
            return None
        parameters = {str(k): float(v) for k, v in parameters.items()}
        if any(not re.fullmatch(r"[a-z][a-z0-9_]{0,20}", k) or k in ("x", *(_FUNCTIONS.keys()))
               or not math.isfinite(v) or abs(v) > 10000 for k, v in parameters.items()):
            return None
        if not 10 <= points <= 200 or not all(math.isfinite(v) and abs(v) <= 10000 for v in (start, stop)) or stop <= start:
            return None
        for formula in (expression, baseline):
            if len(formula) > 180:
                return None
            tree = ast.parse(formula, mode="eval")
            if len(list(ast.walk(tree))) > 60:
                return None
            _safe_formula(tree.body, {"x": 1, **parameters})
    except (KeyError, TypeError, ValueError, SyntaxError, ArithmeticError, OverflowError):
        return None
    return {"expression": expression, "baseline": baseline, "x_start": start,
            "x_stop": stop, "points": points, "parameters": parameters,
            "hypothesis": str(data.get("hypothesis", ""))[:500],
            "limitations": str(data.get("limitations", ""))[:500]}


def _safe_formula(node: ast.AST, names: dict[str, float]) -> float:
    if isinstance(node, ast.Constant) and type(node.value) in (int, float):
        value = float(node.value)
    elif isinstance(node, ast.Name) and node.id in names:
        value = float(names[node.id])
    elif isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
        value = _safe_formula(node.operand, names) * (-1 if isinstance(node.op, ast.USub) else 1)
    elif isinstance(node, ast.BinOp) and type(node.op) in _OPERATORS:
        left, right = _safe_formula(node.left, names), _safe_formula(node.right, names)
        if isinstance(node.op, ast.Pow) and (abs(right) > 6 or abs(left) > 10000):
            raise ValueError("exponent out of range")
        value = _OPERATORS[type(node.op)](left, right)
    elif isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and node.func.id in _FUNCTIONS and not node.keywords and len(node.args) == 1:
        value = _FUNCTIONS[node.func.id](_safe_formula(node.args[0], names))
    else:
        raise ValueError("unsupported expression")
    if not math.isfinite(value) or abs(value) > 1e12:
        raise ValueError("non-finite result")
    return value


def run_simulation(project: Path, cycle: int, protocol: dict) -> dict:
    """Evaluate two bounded formulas on the same grid; output is simulation only."""
    xs = [protocol["x_start"] + (protocol["x_stop"] - protocol["x_start"]) * i / (protocol["points"] - 1)
          for i in range(protocol["points"])]
    rows = []
    for x in xs:
        names = {"x": x, **protocol["parameters"]}
        try:
            target = _safe_formula(ast.parse(protocol["expression"], mode="eval").body, names)
            control = _safe_formula(ast.parse(protocol["baseline"], mode="eval").body, names)
        except (ValueError, ArithmeticError, OverflowError) as exc:
            raise ValueError("시뮬레이션 수식이 범위 안에서 계산되지 않습니다.") from exc
        rows.append((x, target, control, target - control))
    output = project / "datasets" / f"cycle_{cycle:02d}_simulated_data.csv"
    with output.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["evidence", "x", "model_value", "baseline_value", "difference"])
        writer.writerows(("SIMULATED_MODEL", *row) for row in rows)
    result = {"evidence": "SIMULATED_MODEL", "cycle": cycle, "n": len(rows),
              "mean_difference": round(statistics.mean(row[3] for row in rows), 6),
              "max_absolute_difference": round(max(abs(row[3]) for row in rows), 6),
              "protocol": protocol, "limitation": "수식에서 계산한 가상값이며 물리 실측 또는 외부 데이터가 아닙니다."}
    (project / "results" / f"cycle_{cycle:02d}_analysis.json").write_text(
        json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def decision(text: str, allowed: tuple[str, ...], default: str) -> str:
    match = re.search(r"^NEXT\s*[:：]\s*([A-Z_]+)\s*$", text, re.I | re.M)
    value = match.group(1).upper() if match else ""
    return value if value in allowed else default


def extract_cases(text: str, *, limit: int = 12) -> list[dict]:
    match = re.search(r"CASES_JSON\s*[:：]\s*(\[[\s\S]*?\])\s*(?:\n\S|$)", text)
    if not match:
        return []
    try:
        raw = json.loads(match.group(1))
    except (ValueError, TypeError):
        return []
    if not isinstance(raw, list):
        return []
    cases = []
    for item in raw[:limit]:
        if not isinstance(item, dict):
            continue
        prompt = str(item.get("prompt", "")).strip()[:800]
        if not prompt:
            continue
        cases.append({"id": f"case-{len(cases)+1:03d}", "prompt": prompt,
                      "expected": str(item.get("expected", "")).strip()[:300],
                      "source": "AI_GENERATED", "source_note": str(item.get("source_note", "")).strip()[:300]})
    return cases


def slides_from_text(text: str, *, title: str, subtitle: str) -> list[tuple[str, list[str]]]:
    slides = []
    current_title = ""
    bullets: list[str] = []
    for line in text.splitlines():
        match = re.match(r"^\s*(?:#{1,3}\s*)?SLIDE\s*\d*\s*[:：|]\s*(.+?)\s*$", line, re.I)
        if match:
            if current_title:
                slides.append((current_title, bullets[:7]))
            current_title, bullets = match.group(1)[:90], []
        elif current_title and line.strip():
            clean = re.sub(r"^\s*[-*•\d.)]+\s*", "", line).strip()
            if clean:
                bullets.append(clean[:240])
    if current_title:
        slides.append((current_title, bullets[:7]))
    if not slides:
        slides = [("연구 질문", [subtitle]), ("발표 초안", [text[:240] or "내용 생성 실패 — 검토 후 수정 필요"])]
    return [(title, [subtitle])] + slides[:15] + [("검토 사항", ["모델이 작성한 초안입니다. 출처·실험 결과·수치를 사람이 검토해야 합니다."])]


def write_pptx(path: Path, slides: list[tuple[str, list[str]]]) -> None:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    for index, (heading, bullets) in enumerate(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = RGBColor(248, 249, 255)
        stripe = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(.09))
        stripe.fill.solid(); stripe.fill.fore_color.rgb = RGBColor(97, 104, 238)
        stripe.line.fill.background()
        titlebox = slide.shapes.add_textbox(Inches(.85), Inches(.65), Inches(11.65), Inches(1.25))
        p = titlebox.text_frame.paragraphs[0]
        p.text = heading
        p.font.name = "Malgun Gothic"; p.font.size = Pt(34 if index else 40); p.font.bold = True
        p.font.color.rgb = RGBColor(24, 39, 77)
        box = slide.shapes.add_textbox(Inches(.95), Inches(2.0), Inches(11.5), Inches(4.55))
        tf = box.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for pos, bullet in enumerate(bullets):
            para = tf.paragraphs[0] if pos == 0 else tf.add_paragraph()
            para.text = ("• " if index else "") + bullet
            para.font.name = "Malgun Gothic"; para.font.size = Pt(24 if index else 28)
            para.font.color.rgb = RGBColor(45, 58, 91)
            para.space_after = Pt(20)
        foot = slide.shapes.add_textbox(Inches(.95), Inches(7.02), Inches(11.5), Inches(.3))
        para = foot.text_frame.paragraphs[0]
        para.text = f"과학 R&E Lab  •  검토용 초안  •  {index + 1}/{len(slides)}"
        para.font.name = "Malgun Gothic"; para.font.size = Pt(10)
        para.font.color.rgb = RGBColor(106, 119, 149)
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))


def write_presentation(project: Path, stage: str, text: str, topic: str, question: str) -> list[str]:
    directory = project / "presentations" / stage
    directory.mkdir(parents=True, exist_ok=True)
    source = directory / "slides.md"
    source.write_text(text, encoding="utf-8")
    deck = directory / "slides.pptx"
    write_pptx(deck, slides_from_text(text, title=topic or "과학고 R&E 연구", subtitle=question or "연구 질문 검토 중"))
    return [source.relative_to(project).as_posix(), deck.relative_to(project).as_posix()]


def write_report(project: Path, stage: str, text: str, topic: str) -> list[str]:
    directory = project / "presentations" / stage
    directory.mkdir(parents=True, exist_ok=True)
    source = directory / "report.md"
    source.write_text(text, encoding="utf-8")
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(2.2); section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4); section.right_margin = Cm(2.4)
    normal = document.styles["Normal"]
    normal.font.name = "Malgun Gothic"; normal.font.size = DocPt(10.5)
    document.add_heading(topic or "과학고 R&E 연구보고서", level=0)
    document.add_paragraph("과학 R&E Lab · 검토용 초안 · 실제 출처와 실험 결과는 승인 전에 확인 필요")
    for line in text.splitlines():
        clean = line.strip()
        if not clean:
            continue
        if clean.startswith("### "):
            document.add_heading(clean[4:], level=3)
        elif clean.startswith("## "):
            document.add_heading(clean[3:], level=2)
        elif clean.startswith("# "):
            document.add_heading(clean[2:], level=1)
        elif re.match(r"^[-*] ", clean):
            document.add_paragraph(clean[2:], style="List Bullet")
        else:
            document.add_paragraph(clean)
    docx = directory / "report.docx"
    document.save(str(docx))
    return [source.relative_to(project).as_posix(), docx.relative_to(project).as_posix()]


def observed_analysis(project: Path, cycle: int, cases: list[dict], observations: list[dict], *, replicates: int = 1000) -> dict:
    """Summarize only recorded BAZE responses; bootstrap is labelled simulation."""
    by_provider: dict[str, list[dict]] = {}
    replicates = max(100, min(int(replicates), 10000))
    case_by_id = {item["id"]: item for item in cases}
    for row in observations:
        by_provider.setdefault(row["provider"], []).append(row)
    summary = {"evidence": "OBSERVED_BAZE", "cycle": cycle, "case_origin": "AI_GENERATED",
               "cases": len(cases), "observations": len(observations),
               "bootstrap_replicates": replicates, "bootstrap_evidence": "SIMULATED_RESAMPLING", "models": {}}
    rng = random.Random(20261012 + cycle)
    for provider, rows in sorted(by_provider.items()):
        lengths = [len(row["response"]) for row in rows]
        if not lengths:
            continue
        resampled = [statistics.mean(rng.choices(lengths, k=len(lengths))) for _ in range(replicates)]
        resampled.sort()
        exact = [int(row["response"].strip() == case_by_id.get(row["case_id"], {}).get("expected", ""))
                 for row in rows if case_by_id.get(row["case_id"], {}).get("expected")]
        summary["models"][provider] = {
            "n": len(rows), "mean_characters": round(statistics.mean(lengths), 2),
            "bootstrap_95_percent_mean_characters": [round(resampled[int(.025 * (replicates - 1))], 2),
                                                     round(resampled[int(.975 * (replicates - 1))], 2)],
            "exact_match_n": len(exact),
            "exact_match_rate": round(statistics.mean(exact), 4) if exact else None,
            "note": "문자 수는 품질 점수가 아닙니다. 부트스트랩 구간은 관찰값 재표본추출 시뮬레이션입니다."}
    target = project / "results" / f"cycle_{cycle:02d}_analysis.json"
    target.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary
